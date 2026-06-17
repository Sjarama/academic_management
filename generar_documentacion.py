from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

style_normal = doc.styles['Normal']
style_normal.font.name = 'Calibri'
style_normal.font.size = Pt(11)

def set_font(run, bold=False, size=11, color=None, italic=False):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = 'Calibri'
    if color:
        run.font.color.rgb = RGBColor(*color)

def heading(doc, text, level=1, color=(0,0,0)):
    h = doc.add_heading('', level=level)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = h.add_run(text)
    run.font.name = 'Calibri'
    run.font.color.rgb = RGBColor(*color)
    if level == 1:
        run.font.size = Pt(16)
        run.bold = True
    elif level == 2:
        run.font.size = Pt(13)
        run.bold = True
    else:
        run.font.size = Pt(11)
        run.bold = True
    return h

def paragraph(doc, text, indent=0, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if indent:
        p.paragraph_format.left_indent = Inches(indent)
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(11)
    return p

def bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(11)
    return p

def code_block(doc, lines):
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Inches(0.4)
        run = p.add_run(line)
        run.font.name = 'Courier New'
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(30, 30, 30)
    doc.add_paragraph()

def add_table(doc, headers, rows):
    table = doc.add_table(rows=1+len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.paragraphs[0].clear()
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.name = 'Calibri'
        run.font.size = Pt(10)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        shading = OxmlElement('w:shd')
        shading.set(qn('w:val'), 'clear')
        shading.set(qn('w:color'), 'auto')
        shading.set(qn('w:fill'), '2E5FA3')
        cell._tc.get_or_add_tcPr().append(shading)
        run.font.color.rgb = RGBColor(255, 255, 255)
    for ri, row_data in enumerate(rows):
        row = table.rows[ri+1]
        fill = 'F2F2F2' if ri % 2 == 0 else 'FFFFFF'
        for ci, val in enumerate(row_data):
            cell = row.cells[ci]
            cell.paragraphs[0].clear()
            run = cell.paragraphs[0].add_run(str(val))
            run.font.name = 'Calibri'
            run.font.size = Pt(10)
            shading = OxmlElement('w:shd')
            shading.set(qn('w:val'), 'clear')
            shading.set(qn('w:color'), 'auto')
            shading.set(qn('w:fill'), fill)
            cell._tc.get_or_add_tcPr().append(shading)
    doc.add_paragraph()
    return table

for section in doc.sections:
    section.top_margin    = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin   = Cm(3)
    section.right_margin  = Cm(2.5)

# ══════════════════════════════════════════════════════════
# CARÁTULA
# ══════════════════════════════════════════════════════════
doc.add_paragraph()
doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('SISTEMA DE GESTIÓN ACADÉMICA')
run.bold = True; run.font.size = Pt(20); run.font.name = 'Calibri'
run.font.color.rgb = RGBColor(0, 51, 102)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('Documentación Técnica — EP3')
run.font.size = Pt(14); run.font.name = 'Calibri'
run.font.color.rgb = RGBColor(70, 70, 70)

doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('Arquitectura de Microservicios con Spring Boot, React y Docker')
run.italic = True; run.font.size = Pt(12); run.font.name = 'Calibri'

doc.add_paragraph(); doc.add_paragraph(); doc.add_paragraph()

for label, value in [
    ('Autor:',         'Sebastián Jara'),
    ('Asignatura:',    'DSY1106 — Desarrollo Fullstack III'),
    ('Evaluación:',    'EP3'),
    ('Institución:',   'Duoc UC — Sede Mac-Iver'),
    ('Fecha:',         'Junio 2026'),
    ('Repositorio:',   'github.com/Sjarama/academic_management'),
]:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p.add_run(label + '  '); r1.bold = True; r1.font.name = 'Calibri'; r1.font.size = Pt(11)
    r2 = p.add_run(value);         r2.font.name = 'Calibri'; r2.font.size = Pt(11)

doc.add_page_break()

# ══════════════════════════════════════════════════════════
# 1. INTRODUCCIÓN
# ══════════════════════════════════════════════════════════
heading(doc, '1. Introducción', 1, (0, 51, 102))
paragraph(doc,
    'El presente documento describe el diseño, arquitectura y estado de implementación del '
    'Sistema de Gestión Académica, un proyecto desarrollado para la evaluación parcial EP3 '
    'de la asignatura DSY1106 — Desarrollo Fullstack III en Duoc UC.')
paragraph(doc,
    'El sistema aplica una arquitectura de microservicios donde cada dominio del negocio '
    '(usuarios, cursos, pagos, notificaciones) es administrado por un servicio autónomo con '
    'su propia base de datos MySQL, sus propios tests unitarios y su propio contenedor Docker.')
paragraph(doc,
    'A la fecha de este documento, todos los microservicios están completamente implementados '
    'y funcionales, con cobertura JaCoCo superior al 60% en cada uno. El frontend fue '
    'desarrollado en React 18 con Vite y se sirve mediante nginx.')

doc.add_paragraph()

# ══════════════════════════════════════════════════════════
# 2. ARQUITECTURA DEL SISTEMA
# ══════════════════════════════════════════════════════════
heading(doc, '2. Arquitectura del Sistema', 1, (0, 51, 102))

heading(doc, '2.1 Diagrama General', 2, (46, 95, 163))
code_block(doc, [
    '┌─────────────────────────────────────────────────────────────────┐',
    '│              Frontend React (Vite + Chart.js)                   │',
    '│                    http://localhost:5501                        │',
    '└──────────────────────────┬──────────────────────────────────────┘',
    '                           │ HTTP REST / CORS',
    '        ┌──────────────────┼───────────────────┐──────────────────┐',
    '        │                  │                   │                  │',
    '        ▼                  ▼                   ▼                  ▼',
    ' ┌─────────────┐  ┌──────────────┐  ┌─────────────┐  ┌──────────────┐',
    ' │user-service │  │course-service│  │payment-svc  │  │notif-service │',
    ' │   :8081     │  │    :8082     │  │   :8083     │  │   :8084      │',
    ' └──────┬──────┘  └──────┬───────┘  └──────┬──────┘  └──────┬───────┘',
    '        │                │                  │                │',
    '        ▼                ▼                  ▼                ▼',
    '   ┌─────────┐    ┌──────────┐    ┌──────────┐    ┌──────────────┐',
    '   │ user-db │    │course-db │    │payment-db│    │notification-db│',
    '   │  :3307  │    │  :3308   │    │  :3309   │    │    :3310      │',
    '   └─────────┘    └──────────┘    └──────────┘    └──────────────┘',
    '',
    '                      ┌─────────────────┐',
    '                      │    RabbitMQ     │',
    '                      │ :5672 / :15672  │',
    '                      └─────────────────┘',
])

heading(doc, '2.2 Microservicios', 2, (46, 95, 163))
add_table(doc,
    ['Servicio', 'Puerto', 'Base de datos', 'Descripción'],
    [
        ['user-service',         '8081', 'db_usuarios (3307)',      'Estudiantes y profesores'],
        ['course-service',       '8082', 'db_cursos (3308)',        'Cursos académicos'],
        ['payment-service',      '8083', 'db_pagos (3309)',         'Pagos estudiantiles'],
        ['notification-service', '8084', 'db_notificaciones (3310)','Notificaciones del sistema'],
    ]
)

heading(doc, '2.3 Stack Tecnológico', 2, (46, 95, 163))
add_table(doc,
    ['Capa', 'Tecnología', 'Versión'],
    [
        ['Backend',      'Java + Spring Boot',    '21 / 3.x'],
        ['Persistencia', 'MySQL + Spring Data JPA','8.0 / Hibernate'],
        ['Frontend',     'React + Vite',           '18 / 5.x'],
        ['Charts',       'Chart.js + react-chartjs-2','4.4'],
        ['HTTP Client',  'Axios',                  '1.7'],
        ['Contenedores', 'Docker + Docker Compose','Desktop'],
        ['Mensajería',   'RabbitMQ',               '3-management'],
        ['Tests',        'JUnit 5 + Mockito + JaCoCo','incluido en SB'],
    ]
)

doc.add_page_break()

# ══════════════════════════════════════════════════════════
# 3. DESCRIPCIÓN DE PERSISTENCIA
# ══════════════════════════════════════════════════════════
heading(doc, '3. Descripción de Persistencia', 1, (0, 51, 102))

paragraph(doc,
    'Cada microservicio gestiona su propia base de datos MySQL. No existe ninguna tabla compartida '
    'entre servicios. La comunicación entre servicios se realiza exclusivamente por HTTP REST, '
    'nunca por acceso directo a la base de datos de otro servicio.')

paragraph(doc,
    'Hibernate (a través de Spring Data JPA) gestiona el esquema automáticamente con '
    'spring.jpa.hibernate.ddl-auto=update. Al arrancar cada servicio, un ApplicationRunner '
    'verifica si la tabla está vacía y, de estarlo, inserta los datos de prueba.')

doc.add_paragraph()

heading(doc, '3.1 user-service — db_usuarios', 2, (46, 95, 163))

heading(doc, 'Tabla: students', 3)
add_table(doc,
    ['Campo', 'Tipo SQL', 'Restricciones', 'Descripción'],
    [
        ['id',        'BIGINT',       'PK, AUTO_INCREMENT',       'Identificador único del estudiante'],
        ['first_name','VARCHAR(255)',  'NOT NULL',                 'Nombre del estudiante'],
        ['last_name', 'VARCHAR(255)',  'NOT NULL',                 'Apellido del estudiante'],
        ['email',     'VARCHAR(255)',  'NOT NULL, UNIQUE',         'Correo electrónico (único en el sistema)'],
        ['phone',     'VARCHAR(255)',  'NULL',                     'Teléfono de contacto'],
        ['address',   'VARCHAR(255)',  'NULL',                     'Dirección del estudiante'],
        ['comuna',    'VARCHAR(255)',  'NULL',                     'Comuna de residencia'],
    ]
)

heading(doc, 'Tabla: teachers', 3)
add_table(doc,
    ['Campo', 'Tipo SQL', 'Restricciones', 'Descripción'],
    [
        ['id',         'BIGINT',      'PK, AUTO_INCREMENT',       'Identificador único del profesor'],
        ['first_name', 'VARCHAR(255)', 'NOT NULL',                'Nombre del profesor'],
        ['last_name',  'VARCHAR(255)', 'NOT NULL',                'Apellido del profesor'],
        ['email',      'VARCHAR(255)', 'NOT NULL, UNIQUE',        'Correo electrónico'],
        ['phone',      'VARCHAR(255)', 'NULL',                    'Teléfono de contacto'],
        ['address',    'VARCHAR(255)', 'NULL',                    'Dirección'],
        ['specialty',  'VARCHAR(255)', 'NULL',                    'Especialidad o área de enseñanza'],
    ]
)

paragraph(doc, 'Datos iniciales: 20 profesores y 400 estudiantes generados automáticamente al arrancar.')
doc.add_paragraph()

heading(doc, '3.2 course-service — db_cursos', 2, (46, 95, 163))

heading(doc, 'Tabla: courses', 3)
add_table(doc,
    ['Campo', 'Tipo SQL', 'Restricciones', 'Descripción'],
    [
        ['id',                 'BIGINT',  'PK, AUTO_INCREMENT',   'Identificador único del curso'],
        ['name',               'VARCHAR(255)', 'NOT NULL',         'Nombre del curso'],
        ['description',        'TEXT',    'NULL',                  'Descripción detallada'],
        ['teacher_id',         'BIGINT',  'NOT NULL',              'ID del profesor responsable (referencia a user-service)'],
        ['credits',            'INT',     'NULL',                  'Número de créditos del curso'],
        ['max_students',       'INT',     'NULL',                  'Cupo máximo de estudiantes'],
        ['approval_percentage','INT',     'CHECK(1–100)',          'Porcentaje de aprobación requerido'],
    ]
)

paragraph(doc,
    'Nota: teacher_id no es una foreign key en la base de datos porque el profesor vive en una '
    'base de datos diferente (db_usuarios). La integridad referencial se mantiene a nivel de '
    'aplicación, no a nivel de SQL.')
paragraph(doc, 'Datos iniciales: 20 cursos generados automáticamente.')
doc.add_paragraph()

heading(doc, '3.3 payment-service — db_pagos', 2, (46, 95, 163))

heading(doc, 'Tabla: payments', 3)
add_table(doc,
    ['Campo', 'Tipo SQL', 'Restricciones', 'Descripción'],
    [
        ['id',         'BIGINT',        'PK, AUTO_INCREMENT',  'Identificador único del pago'],
        ['student_id', 'BIGINT',        'NOT NULL',            'ID del estudiante (ref. a user-service)'],
        ['course_id',  'BIGINT',        'NULL',                'ID del curso asociado (opcional)'],
        ['amount',     'DECIMAL(10,2)', 'NOT NULL',            'Monto del pago en pesos chilenos'],
        ['status',     'VARCHAR(255)',  'NOT NULL',            'Estado: Pendiente / Pagado / Vencido'],
        ['due_date',   'DATE',          'NOT NULL',            'Fecha de vencimiento del pago'],
    ]
)

paragraph(doc, 'Datos iniciales: 400 pagos generados automáticamente (uno por estudiante).')
doc.add_paragraph()

heading(doc, '3.4 notification-service — db_notificaciones', 2, (46, 95, 163))

heading(doc, 'Tabla: notifications', 3)
add_table(doc,
    ['Campo', 'Tipo SQL', 'Restricciones', 'Descripción'],
    [
        ['id',         'BIGINT',        'PK, AUTO_INCREMENT',  'Identificador único'],
        ['recipient',  'VARCHAR(255)',  'NOT NULL',            'Email o ID del destinatario'],
        ['message',    'VARCHAR(1000)', 'NOT NULL',            'Contenido de la notificación'],
        ['type',       'VARCHAR(255)',  'NOT NULL',            'Tipo: PAGO_PENDIENTE, BIENVENIDA, etc.'],
        ['sent',       'TINYINT(1)',    'NOT NULL, DEFAULT 0', 'Indica si la notificación fue enviada'],
        ['sent_at',    'DATETIME',      'NULL',                'Timestamp del envío efectivo'],
    ]
)

paragraph(doc, 'Datos iniciales: 200 notificaciones generadas automáticamente.')
doc.add_paragraph()

doc.add_page_break()

# ══════════════════════════════════════════════════════════
# 4. COBERTURA DE TESTS (JaCoCo)
# ══════════════════════════════════════════════════════════
heading(doc, '4. Cobertura de Tests — JaCoCo', 1, (0, 51, 102))

paragraph(doc,
    'Todos los microservicios cuentan con tests unitarios implementados con JUnit 5 y Mockito. '
    'La cobertura fue medida con el plugin JaCoCo de Maven. El requisito mínimo de la evaluación '
    'es 60% de cobertura de instrucciones.')

add_table(doc,
    ['Microservicio', 'Instrucciones cubiertas', 'Instrucciones totales', 'Cobertura %', 'Estado'],
    [
        ['user-service',         '725',  '1105', '65.6%', '✓ Aprobado'],
        ['course-service',       '629',  '730',  '86.2%', '✓ Aprobado'],
        ['payment-service',      '327',  '333',  '98.2%', '✓ Aprobado'],
        ['notification-service', '192',  '280',  '68.6%', '✓ Aprobado'],
    ]
)

paragraph(doc,
    'Los reportes HTML detallados están disponibles en cada microservicio en la ruta: '
    'target/site/jacoco/index.html (generados con ./mvnw test).')

paragraph(doc,
    'La estrategia de testing usa principalmente @ExtendWith(MockitoExtension.class) con '
    '@Mock e @InjectMocks para aislar cada capa del sistema. Para los tests de contexto '
    '(@SpringBootTest), se usa H2 como base de datos en memoria para no requerir MySQL al correr tests.')

doc.add_paragraph()

heading(doc, '4.1 Clases de test por servicio', 2, (46, 95, 163))
add_table(doc,
    ['Servicio', 'Clase de test', 'Qué cubre'],
    [
        ['user-service',         'StudentServiceTest',          'Lógica CRUD del servicio de estudiantes'],
        ['user-service',         'TeacherServiceTest',          'Lógica CRUD del servicio de profesores'],
        ['course-service',       'CourseServiceTest',           'Lógica CRUD del servicio de cursos'],
        ['course-service',       'CourseControllerTest',        'Endpoints REST del controller de cursos'],
        ['course-service',       'ReportControllerTest',        'Generación de reporte HTML'],
        ['payment-service',      'PaymentControllerTest',       'Endpoints REST del controller de pagos'],
        ['notification-service', 'NotificationServiceTest',     'Lógica del servicio de notificaciones'],
        ['notification-service', 'NotificationControllerTest',  'Endpoints REST del controller'],
    ]
)

doc.add_page_break()

# ══════════════════════════════════════════════════════════
# 5. API ENDPOINTS
# ══════════════════════════════════════════════════════════
heading(doc, '5. API Endpoints', 1, (0, 51, 102))

paragraph(doc,
    'La colección Postman completa está disponible en postman/Academic_Management_API.postman_collection.json. '
    'A continuación se listan todos los endpoints por servicio.')

heading(doc, '5.1 user-service (puerto 8081)', 2, (46, 95, 163))
add_table(doc,
    ['Método', 'Endpoint', 'Descripción'],
    [
        ['GET',    '/api/students',           'Listar todos los estudiantes'],
        ['GET',    '/api/students/{id}',      'Obtener estudiante por ID'],
        ['GET',    '/api/students/comunas',   'Listar comunas disponibles'],
        ['POST',   '/api/students',           'Crear nuevo estudiante (409 si email duplicado)'],
        ['PUT',    '/api/students/{id}',      'Actualizar estudiante'],
        ['DELETE', '/api/students/{id}',      'Eliminar estudiante'],
        ['GET',    '/api/teachers',           'Listar todos los profesores'],
        ['GET',    '/api/teachers/{id}',      'Obtener profesor por ID'],
        ['POST',   '/api/teachers',           'Crear nuevo profesor'],
        ['PUT',    '/api/teachers/{id}',      'Actualizar profesor'],
        ['DELETE', '/api/teachers/{id}',      'Eliminar profesor'],
    ]
)

heading(doc, '5.2 course-service (puerto 8082)', 2, (46, 95, 163))
add_table(doc,
    ['Método', 'Endpoint', 'Descripción'],
    [
        ['GET',    '/api/courses',                     'Listar todos los cursos'],
        ['GET',    '/api/courses/{id}',                'Obtener curso por ID'],
        ['GET',    '/api/courses/teacher/{teacherId}', 'Cursos de un profesor'],
        ['POST',   '/api/courses',                     'Crear curso'],
        ['PUT',    '/api/courses/{id}',                'Actualizar curso'],
        ['DELETE', '/api/courses/{id}',                'Eliminar curso'],
        ['GET',    '/report',                          'Reporte HTML (cursos + estudiantes)'],
    ]
)

heading(doc, '5.3 payment-service (puerto 8083)', 2, (46, 95, 163))
add_table(doc,
    ['Método', 'Endpoint', 'Descripción'],
    [
        ['GET',    '/api/payments',       'Listar todos los pagos'],
        ['GET',    '/api/payments/{id}',  'Obtener pago por ID'],
        ['POST',   '/api/payments',       'Registrar nuevo pago'],
        ['PUT',    '/api/payments/{id}',  'Actualizar pago (ej: marcar como pagado)'],
        ['DELETE', '/api/payments/{id}',  'Eliminar pago'],
    ]
)

heading(doc, '5.4 notification-service (puerto 8084)', 2, (46, 95, 163))
add_table(doc,
    ['Método', 'Endpoint', 'Descripción'],
    [
        ['GET',    '/api/notifications',       'Listar todas las notificaciones'],
        ['GET',    '/api/notifications/{id}',  'Obtener notificación por ID'],
        ['POST',   '/api/notifications',       'Crear notificación'],
        ['PUT',    '/api/notifications/{id}',  'Actualizar notificación'],
        ['DELETE', '/api/notifications/{id}',  'Eliminar notificación'],
    ]
)

doc.add_page_break()

# ══════════════════════════════════════════════════════════
# 6. FRONTEND REACT
# ══════════════════════════════════════════════════════════
heading(doc, '6. Frontend React', 1, (0, 51, 102))

paragraph(doc,
    'El frontend fue desarrollado en React 18 con Vite como bundler. Se sirve en producción '
    'mediante un contenedor nginx (Docker multi-stage build). Disponible en http://localhost:5501.')

heading(doc, '6.1 Estructura de componentes', 2, (46, 95, 163))
add_table(doc,
    ['Archivo', 'Descripción'],
    [
        ['App.jsx',                  'Componente raíz. Maneja navegación por estado y refs para Topbar→Modal.'],
        ['components/Sidebar.jsx',   'Navegación lateral oscura con 6 secciones.'],
        ['components/Topbar.jsx',    'Barra superior con botones Actualizar / Exportar / Nuevo.'],
        ['components/DataTable.jsx', 'Tabla genérica con badges de estado, edición y eliminación.'],
        ['components/EntityModal.jsx','Modal dinámico para crear y editar entidades (soporta select, textarea, checkbox, date).'],
        ['components/ToastContainer.jsx','Notificaciones toast success/error/info con auto-dismiss.'],
        ['pages/Dashboard.jsx',      'Métricas en tiempo real + Doughnut (comunas) + Bar (reprobación) con Chart.js.'],
        ['pages/EntityPage.jsx',     'Página CRUD genérica con buscador en tiempo real y exportación a Excel.'],
        ['config/services.js',       'Configuración centralizada: URL, columnas, labels y campos de formulario.'],
        ['api/client.js',            'Funciones fetchAll, createEntity, updateEntity, deleteEntity con Axios.'],
        ['hooks/useToast.js',        'Custom Hook para gestión de toasts (estado + auto-cleanup).'],
    ]
)

heading(doc, '6.2 Funcionalidades implementadas', 2, (46, 95, 163))
bullet(doc, 'CRUD completo para: Estudiantes, Profesores, Cursos, Pagos, Notificaciones.')
bullet(doc, 'Dashboard con 4 métricas (estudiantes, cursos, pagos pendientes, estudiantes al día).')
bullet(doc, 'Gráfico de torta (Doughnut): distribución de estudiantes por comuna.')
bullet(doc, 'Gráfico de barras (Bar): cursos con mayor porcentaje de reprobación (rojo/amarillo/verde).')
bullet(doc, 'Buscador en tiempo real en cada tabla (filtro client-side).')
bullet(doc, 'Exportación de datos a Excel (.xls) desde el botón Exportar.')
bullet(doc, 'Toast notifications para feedback de operaciones.')
bullet(doc, 'Modal dinámico con campos según el tipo de entidad.')

doc.add_page_break()

# ══════════════════════════════════════════════════════════
# 7. DESPLIEGUE CON DOCKER
# ══════════════════════════════════════════════════════════
heading(doc, '7. Despliegue con Docker', 1, (0, 51, 102))

heading(doc, '7.1 Contenedores del sistema', 2, (46, 95, 163))
add_table(doc,
    ['Contenedor', 'Imagen base', 'Puerto host', 'Puerto interno'],
    [
        ['user-db',              'mysql:8.0',              '3307',       '3306'],
        ['user-service',         'Maven + eclipse-temurin','8081',       '8081'],
        ['course-db',            'mysql:8.0',              '3308',       '3306'],
        ['course-service',       'Maven + eclipse-temurin','8082',       '8082'],
        ['payment-db',           'mysql:8.0',              '3309',       '3306'],
        ['payment-service',      'Maven + eclipse-temurin','8083',       '8083'],
        ['notification-db',      'mysql:8.0',              '3310',       '3306'],
        ['notification-service', 'Maven + eclipse-temurin','8084',       '8084'],
        ['frontend',             'node:20 + nginx:alpine', '5501',       '80'],
        ['rabbitmq',             'rabbitmq:3-management',  '5672 / 15672','5672 / 15672'],
    ]
)

heading(doc, '7.2 Comandos principales', 2, (46, 95, 163))
code_block(doc, [
    '# Levantar todo el sistema',
    'docker-compose up --build -d',
    '',
    '# Ver estado de contenedores',
    'docker ps',
    '',
    '# Ver logs de un servicio',
    'docker logs user-service -f',
    '',
    '# Correr tests con JaCoCo',
    'cd user-service && ./mvnw test',
    '# Reporte en: target/site/jacoco/index.html',
])

doc.add_page_break()

# ══════════════════════════════════════════════════════════
# 8. ESTADO EP3
# ══════════════════════════════════════════════════════════
heading(doc, '8. Estado del Proyecto — EP3', 1, (0, 51, 102))

add_table(doc,
    ['Requisito EP3', 'Estado', 'Observación'],
    [
        ['4 microservicios funcionales',       'Completado', 'user, course, payment, notification — CRUD completo'],
        ['Base de datos por microservicio',    'Completado', 'MySQL 8.0 independiente por servicio'],
        ['Docker Compose orquestado',          'Completado', 'docker-compose up --build -d'],
        ['Frontend con framework moderno',     'Completado', 'React 18 + Vite + Chart.js'],
        ['package.json (NPM)',                 'Completado', 'frontend-react/package.json'],
        ['Tests unitarios JaCoCo ≥ 60%',      'Completado', 'user 65.6% | course 86.2% | payment 98.2% | notif 68.6%'],
        ['README por microservicio',           'Completado', 'README.md en cada carpeta de servicio'],
        ['Colección Postman',                  'Completado', 'postman/Academic_Management_API.postman_collection.json'],
        ['Documentación técnica',              'Completado', 'Este documento (Documentacion_Sistema_Gestion_Academica.docx)'],
        ['repositorios.txt',                   'Completado', 'Archivo con links a GitHub'],
    ]
)

doc.add_paragraph()

heading(doc, '9. Conclusiones', 1, (0, 51, 102))
paragraph(doc,
    'El Sistema de Gestión Académica implementa una arquitectura de microservicios completa y '
    'funcional. Cada uno de los cuatro servicios está completamente implementado, con CRUD, '
    'datos de prueba, tests unitarios con cobertura superior al 60%, Dockerfile y README propios.')
paragraph(doc,
    'El frontend en React consume todos los microservicios y presenta dashboards con gráficos '
    'interactivos, tablas con búsqueda en tiempo real y modales para gestión de datos. '
    'Todos los componentes se orquestan mediante Docker Compose en un solo comando.')
paragraph(doc,
    'Este proyecto demuestra la aplicación práctica de patrones de diseño, separación de '
    'responsabilidades y despliegue contenerizado en un sistema de software real.')

# ══════════════════════════════════════════════════════════
# GUARDAR
# ══════════════════════════════════════════════════════════
output_path = r'c:\Users\Sebs\Documents\GitHub\academic_management\Documentacion_Sistema_Gestion_Academica.docx'
doc.save(output_path)
print(f'Documento guardado en: {output_path}')
